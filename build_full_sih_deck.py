import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Styling Constants
COLOR_BG = RGBColor(10, 14, 23)          # Dark space navy
COLOR_CARD = RGBColor(18, 25, 41)        # Glass card slate
COLOR_PRIMARY = RGBColor(56, 189, 248)   # Electric Cyan
COLOR_ACCENT = RGBColor(244, 63, 94)     # Warning Rose
COLOR_AMBER = RGBColor(245, 158, 11)     # Radar Amber
COLOR_EMERALD = RGBColor(16, 185, 129)   # Live Emerald
COLOR_TEXT_MAIN = RGBColor(248, 250, 252)# Bright White
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)# Slate Gray
COLOR_BORDER = RGBColor(30, 41, 59)      # Border

HERO_IMG = r"C:\Users\LENOVO\.gemini\antigravity\brain\2eef33c0-447e-432e-8c02-c6cf86f11bda\nowcast_gis_dashboard_1790699495598.jpg"

def create_slide_base(prs, category_tag, slide_title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.fill.background()
    
    # Header Bar Box
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.9))
    tf = tb.text_frame
    tf.word_wrap = True
    
    p0 = tf.paragraphs[0]
    p0.text = category_tag.upper()
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_PRIMARY
    
    p1 = tf.add_paragraph()
    p1.text = slide_title
    p1.font.size = Pt(22)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_TEXT_MAIN
    p1.space_before = Pt(3)
    
    return slide

def style_cell(cell, text, bold=False, font_size=10, color=COLOR_TEXT_MAIN, bg_color=None):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color

# ==========================================
# SLIDE 1 — IDEA TITLE (Cover Slide)
# ==========================================
s1 = prs.slides.add_slide(prs.slide_layouts[6])
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg1.fill.solid()
bg1.fill.fore_color.rgb = COLOR_BG
bg1.line.fill.background()

tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(7.2), Inches(5.5))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "SMART INDIA HACKATHON 2026"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

p = tf1.add_paragraph()
p.text = "Nowcast-GIS"
p.font.size = Pt(46)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_MAIN
p.space_before = Pt(6)

p = tf1.add_paragraph()
p.text = "High-Resolution Convective-Scale Weather Prediction & Real-Time GIS Early Warning Engine"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = COLOR_AMBER
p.space_before = Pt(10)

p = tf1.add_paragraph()
p.text = "Bridging the 0–6 Hour Weather Gap: From Satellite Thermal Detection to Ground-Level Evacuation."
p.font.size = Pt(12)
p.font.italic = True
p.font.color.rgb = COLOR_TEXT_MUTED
p.space_before = Pt(10)

p = tf1.add_paragraph()
p.text = "An end-to-end, AI-powered convective nowcasting platform that ingests INSAT-3D thermal IR and Doppler radar to detect storm updrafts 30–45 minutes before rain begins and projects severe storm advection up to 6 hours with real-time GIS countdown clocks."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MAIN
p.space_before = Pt(14)

p = tf1.add_paragraph()
p.text = "\nTeam: [Your Team Name]  |  Problem Statement: Convective-Scale Severe Weather Nowcasting"
p.font.size = Pt(11)
p.font.color.rgb = COLOR_PRIMARY

if os.path.exists(HERO_IMG):
    s1.shapes.add_picture(HERO_IMG, Inches(7.6), Inches(1.1), Inches(5.1), Inches(5.2))


# ==========================================
# SLIDE 2 — PROPOSED SOLUTION
# ==========================================
s2 = create_slide_base(prs, "SLIDE 2 — PROPOSED SOLUTION", "The Nowcasting Problem & Our Integrated GIS Solution")

# Left Column (The Problem & 3-Step Workflow)
tb2_left = s2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.8), Inches(5.4))
tf2_left = tb2_left.text_frame
tf2_left.word_wrap = True

p = tf2_left.paragraphs[0]
p.text = "WHAT IS THE PROBLEM?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

p = tf2_left.add_paragraph()
p.text = "• Severe thunderstorms & lightning claim 2,500+ Indian lives annually.\n• Traditional NWP models have 3–6 hour compute latency and miss rapid convective storm births.\n• Disaster authorities (NDRF/SDMA) receive zero sub-hour tactical lead time."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MUTED
p.space_before = Pt(4)

p = tf2_left.add_paragraph()
p.text = "\nHOW IT WORKS (3 SIMPLE STEPS):"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.space_before = Pt(8)

p = tf2_left.add_paragraph()
p.text = "1. Ingest & Detect: Monitors INSAT-3D cloud-top cooling rates (ΔT/Δt) to spot storm updrafts 30–45 mins before radar echoes form.\n2. Nowcast (0–6 Hours): Dual-engine advection (OpenCV Optical Flow + ConvLSTM) outputs 12 forecast frames at 30-min steps.\n3. Warn & Act: Clusters severe echoes (≥50 dBZ) and lightning into active hazard alert zones with automated, ticking countdown clocks."
p.font.size = Pt(11)
p.font.color.rgb = COLOR_TEXT_MAIN
p.space_before = Pt(4)

# Right Column (Unique Features)
tb2_right = s2.shapes.add_textbox(Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.4))
tf2_right = tb2_right.text_frame
tf2_right.word_wrap = True

p = tf2_right.paragraphs[0]
p.text = "WHAT MAKES NOWCAST-GIS UNIQUE:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

features = [
    ("Pre-Radar Convective Initiation", "Identifies storm updrafts 30–45 mins before ground rain via satellite thermal cooling."),
    ("Live Arrival Countdown Clocks", "Zone cards calculate exact ticking countdowns (MM:SS) to storm arrival."),
    ("6-Hour Interactive Timeline", "13 time capsules (T+0 to T+360m) with animated Play/Pause playback."),
    ("Dual-Model Engine", "Combines OpenCV Farneback Optical Flow with deep learning ConvLSTM."),
    ("Multi-Hazard Data Fusion", "Unified GeoJSON payload merging radar dBZ polygons, CI points & lightning."),
    ("Offline Demo / Replay Mode", "Embedded historical replay toggle for venue demos without Wi-Fi."),
    ("WebSocket Real-Time Stream", "Notifies clients within 3 seconds of new file ingestion.")
]

for title, desc in features:
    p = tf2_right.add_paragraph()
    p.text = f"• {title}: {desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(4)


# ==========================================
# SLIDE 3 — TECHNICAL APPROACH
# ==========================================
s3 = create_slide_base(prs, "SLIDE 3 — TECHNICAL APPROACH", "System Architecture & Production Technology Stack")

# Tech Stack Table
table_shape = s3.shapes.add_table(7, 2, Inches(0.8), Inches(1.5), Inches(5.4), Inches(5.3))
t = table_shape.table
t.columns[0].width = Inches(1.8)
t.columns[1].width = Inches(3.6)

tech_data = [
    ("LAYER", "TECHNOLOGIES EMPLOYED"),
    ("GIS Dashboard", "HTML5, Tailwind CSS, Vanilla JS, Leaflet.js 1.9, CartoDB Dark"),
    ("Backend API", "Python 3.14, FastAPI, Uvicorn (ASGI), Pydantic v2"),
    ("Nowcast Models", "OpenCV (Farneback Optical Flow), PyTorch (ConvLSTM)"),
    ("Geospatial Stack", "Rasterio 1.5.1, GDAL/Affine, NumPy 2.5, Pillow"),
    ("CI Detection ML", "Scikit-Learn, LightGBM, Joblib (ΔT/Δt + Texture Var)"),
    ("Real-Time Push", "Native WebSockets (/live) with async 3s file-watcher"),
]

for row_idx, (layer, tech) in enumerate(tech_data):
    is_head = row_idx == 0
    bg_col = COLOR_CARD if not is_head else RGBColor(30, 41, 59)
    style_cell(t.cell(row_idx, 0), layer, bold=True, font_size=10, color=COLOR_PRIMARY if is_head else COLOR_AMBER, bg_color=bg_col)
    style_cell(t.cell(row_idx, 1), tech, bold=is_head, font_size=10, color=COLOR_TEXT_MAIN, bg_color=bg_col)

# Architecture Box on Right
tb3_arch = s3.shapes.add_textbox(Inches(6.5), Inches(1.5), Inches(6.0), Inches(5.3))
tf3_arch = tb3_arch.text_frame
tf3_arch.word_wrap = True

p = tf3_arch.paragraphs[0]
p.text = "END-TO-END PIPELINE ARCHITECTURE:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

arch_steps = [
    ("1. Ingestion Layer", "Pulls INSAT-3D thermal IR, DWR Doppler radar (0.027° grid), and synthetic lightning proxy data into /engine/data/."),
    ("2. Convective Initiation (CI)", "ML classifier extracts cooling rate and texture variance to output geo-referenced CI detection points."),
    ("3. AI Nowcasting Engine", "OpenCV Farneback dense optical flow advects radar reflectivity forward into 12 future frames (30m intervals, 6h horizon)."),
    ("4. Data Readers (readers.py)", "Extracts GeoJSON contour polygons (Moderate 30-40, Heavy 40-50, Severe ≥50 dBZ) and synthesizes warning zones (HZ-01..06)."),
    ("5. FastAPI Service (main.py)", "Exposes /status, /hazards, /forecast, /replay, and WS /live. Verified with 10/10 passing pytest test cases."),
    ("6. Full-Screen Dashboard", "Full-screen Leaflet dark map centered on India with active storm countdown clocks, bottom timeline scrubber, and offline replay mode.")
]

for step_title, step_desc in arch_steps:
    p = tf3_arch.add_paragraph()
    p.text = f"{step_title}: {step_desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(5)


# ==========================================
# SLIDE 4 — FEASIBILITY AND VIABILITY
# ==========================================
s4 = create_slide_base(prs, "SLIDE 4 — FEASIBILITY AND VIABILITY", "Deployment Viability & Engineering Risk Mitigation")

tb4_left = s4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.4), Inches(5.3))
tf4_left = tb4_left.text_frame
tf4_left.word_wrap = True

p = tf4_left.paragraphs[0]
p.text = "WHY THIS IS FEASIBLE RIGHT NOW:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

feasibility = [
    ("Technical Feasibility", "Built on open standards (GeoTIFF, GeoJSON, EPSG:4326). Farneback optical flow executes in < 2 seconds on CPU. 100% operational working prototype."),
    ("Operational Feasibility", "Direct drop-in for NDRF, SDMA, and airport weather offices. No specialized hardware required; runs on standard edge servers and loads on any browser."),
    ("Financial Feasibility", "Zero proprietary software license fees. 100% open-source stack. Negligible deployment cost vs. billions saved in avoided disaster damage.")
]

for title, desc in feasibility:
    p = tf4_left.add_paragraph()
    p.text = f"• {title}:\n  {desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(8)

# Risk & Mitigation Table on Right
table4 = s4.shapes.add_table(5, 3, Inches(6.5), Inches(1.5), Inches(6.0), Inches(5.3))
t4 = table4.table
t4.columns[0].width = Inches(1.8)
t4.columns[1].width = Inches(1.4)
t4.columns[2].width = Inches(2.8)

risk_data = [
    ("CHALLENGE", "RISK", "MITIGATION STRATEGY"),
    ("Radar beam blockage in hills", "Coverage gaps", "Fused satellite thermal IR (INSAT) provides continuous top-down coverage."),
    ("Non-linear storm dissipation", "Advection inaccuracy", "ConvLSTM deep model predicts cell decay; CI detector tracks new births."),
    ("Unreliable field Wi-Fi", "Live demo lag", "Embedded Replay Mode loops pre-baked 771-hazard snapshot completely offline."),
    ("Large GeoJSON payload", "Browser map stutter", "Polygon simplification & 4-decimal truncation cuts payload by >50%."),
]

for row_idx, (c1, c2, c3) in enumerate(risk_data):
    is_head = row_idx == 0
    bg_col = COLOR_CARD if not is_head else RGBColor(30, 41, 59)
    style_cell(t4.cell(row_idx, 0), c1, bold=True, font_size=9, color=COLOR_PRIMARY if is_head else COLOR_AMBER, bg_color=bg_col)
    style_cell(t4.cell(row_idx, 1), c2, bold=is_head, font_size=9, color=COLOR_ACCENT if not is_head else COLOR_PRIMARY, bg_color=bg_col)
    style_cell(t4.cell(row_idx, 2), c3, bold=is_head, font_size=9, color=COLOR_TEXT_MAIN, bg_color=bg_col)


# ==========================================
# SLIDE 5 — IMPACT AND BENEFITS
# ==========================================
s5 = create_slide_base(prs, "SLIDE 5 — IMPACT AND BENEFITS", "Quantifiable Impact & Multi-Sector Beneficiaries")

# Impact Table
table5 = s5.shapes.add_table(6, 3, Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.3))
t5 = table5.table
t5.columns[0].width = Inches(2.2)
t5.columns[1].width = Inches(1.8)
t5.columns[2].width = Inches(2.0)

impact_data = [
    ("PERFORMANCE METRIC", "TRADITIONAL NWP", "NOWCAST-GIS PLATFORM"),
    ("Forecast Update Frequency", "Every 3 to 6 hours", "Every 5 to 15 minutes"),
    ("Spatial Resolution", "12 km – 25 km grid", "0.027° (~3 km) localized"),
    ("Convective Initiation Lead Time", "0 mins (post-rainfall)", "30–45 mins advance detection"),
    ("Hazard Arrival Precision", "District-wide vague alert", "Zone-specific countdown (MM:SS)"),
    ("Client Payload Size", "Gigabytes (binary grids)", "Lightweight GeoJSON (<500 KB)"),
]

for row_idx, (m, old, new) in enumerate(impact_data):
    is_head = row_idx == 0
    bg_col = COLOR_CARD if not is_head else RGBColor(30, 41, 59)
    style_cell(t5.cell(row_idx, 0), m, bold=True, font_size=9, color=COLOR_PRIMARY if is_head else COLOR_TEXT_MAIN, bg_color=bg_col)
    style_cell(t5.cell(row_idx, 1), old, bold=is_head, font_size=9, color=COLOR_TEXT_MUTED if not is_head else COLOR_PRIMARY, bg_color=bg_col)
    style_cell(t5.cell(row_idx, 2), new, bold=True, font_size=9, color=COLOR_EMERALD if not is_head else COLOR_PRIMARY, bg_color=bg_col)

# Beneficiaries Box on Right
tb5_right = s5.shapes.add_textbox(Inches(7.1), Inches(1.5), Inches(5.4), Inches(5.3))
tf5_right = tb5_right.text_frame
tf5_right.word_wrap = True

p = tf5_right.paragraphs[0]
p.text = "TARGET BENEFICIARIES & BROADER IMPACT:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

beneficiaries = [
    ("Disaster Authorities (SDMA / NDRF)", "Increases actionable lead time from 0 to 45 mins. Live countdowns enable targeted siren alerts and localized evacuations."),
    ("Aviation & Airports", "Enables tactical terminal airspace rerouting around severe cells (≥40 dBZ) up to 2 hours ahead; prevents lightning ramp hazards."),
    ("Agriculture & Rural Communities", "Provides farmers advance notice for localized squalls and hailstorms; drastically cuts lightning fatalities in open fields."),
    ("Smart Cities & Municipalities", "Preemptively triggers storm pump stations 45 minutes before cloudbursts hit; diverts urban underpass traffic.")
]

for b_title, b_desc in beneficiaries:
    p = tf5_right.add_paragraph()
    p.text = f"• {b_title}:\n  {b_desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(6)


# ==========================================
# SLIDE 6 — RESEARCH AND REFERENCES
# ==========================================
s6 = create_slide_base(prs, "SLIDE 6 — RESEARCH AND REFERENCES", "Scientific Literature, Technical Standards & Live Prototype")

tb6_left = s6.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.3))
tf6_left = tb6_left.text_frame
tf6_left.word_wrap = True

p = tf6_left.paragraphs[0]
p.text = "ATMOSPHERIC SCIENCE & DOMAIN REFERENCES:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY

refs = [
    ("India Meteorological Department (IMD)", "Operational Doppler Weather Radar (DWR) Nowcasting Guidelines & Reflectivity dBZ thresholds."),
    ("ISRO / MOSDAC", "INSAT-3D/3DR Meteorological Data Products & Thermal IR Specifications."),
    ("World Meteorological Organization (WMO)", "WMO Publication No. 1198 — Guidelines for Nowcasting Techniques."),
    ("Farneback, G. (2003)", "'Two-Frame Motion Estimation Based on Polynomial Expansion' (Foundational basis for optical flow advection)."),
    ("Shi et al. (2015 - NeurIPS)", "'Convolutional LSTM Network: A Machine Learning Approach for Precipitation Nowcasting'.")
]

for title, desc in refs:
    p = tf6_left.add_paragraph()
    p.text = f"• {title}: {desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(5)

tb6_right = s6.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.3))
tf6_right = tb6_right.text_frame
tf6_right.word_wrap = True

p = tf6_right.paragraphs[0]
p.text = "PROJECT VERIFICATION & LIVE PROTOTYPE:"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = COLOR_EMERALD

prototype_items = [
    ("Test Suite Status", "10/10 automated tests passing (pytest engine/api/test_api.py -v)."),
    ("Hazard Coverage", "771 combined active hazards across India grid (200 CI points, 115 lightning strikes, 450 radar polygons, 6 alert zones)."),
    ("GitHub Repository", "Pull Request #5 open into main: Phase 5-6: backend API + GIS dashboard."),
    ("Live Local Dashboard", "http://127.0.0.1:8000/dashboard/ (Full-screen Leaflet map, countdown timers, 6h timeline scrubber)."),
    ("SIH 2026 Reference", "Problem Statement: Convective-Scale Severe Weather Nowcasting Engine.")
]

for title, desc in prototype_items:
    p = tf6_right.add_paragraph()
    p.text = f"• {title}: {desc}"
    p.font.size = Pt(10)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(6)

# Save
output_file = "SIH_2026_Nowcast_GIS.pptx"
prs.save(output_file)
print(f"Presentation successfully created at: {output_file}")
